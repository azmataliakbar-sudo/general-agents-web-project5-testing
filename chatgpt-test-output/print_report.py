from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import os

# Create new document
doc = Document()

# Set page margins
sec = doc.sections[0]
sec.top_margin = Inches(0.7)
sec.bottom_margin = Inches(0.7)
sec.left_margin = Inches(0.8)
sec.right_margin = Inches(0.8)

# Set default font
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(10.5)

# ==================== TITLE ====================
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("Project 5: One Task, Two Harnesses — Updated Full Record")
r.bold = True
r.font.size = Pt(18)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = subtitle.add_run("Claude Code vs VS Code GPT Extension + Plain ChatGPT Manual Reference")
r.bold = True
r.font.size = Pt(14)

doc.add_paragraph()

# ==================== SECTION 1 ====================
doc.add_heading("1. What The Project Was Asking Us To Build", level=1)
doc.add_paragraph(
    "Give the exact same low-stakes task to two different AI harnesses, and observe—not assume—how each one actually handles it. "
    "Map what is really observed to six parts: Heartbeat, Reach, Loop, Spine, Gate, Body. "
    "The assignment marked Cowork and ChatGPT Work as optional, so the official two-harness study used tools actually available: "
    "Claude Code and the VS Code GPT extension."
)

# ==================== SECTION 2 ====================
doc.add_heading("2. Evidence Status — Official Comparison + Additional Reference", level=1)

doc.add_paragraph(
    "Official Harness A vs Harness B comparison: Claude Code vs the VS Code GPT extension (vscode-chat-gpt). "
    "This remains the formal two-harness comparison required by the assignment."
)

doc.add_paragraph(
    "Additional reference test: Plain ChatGPT chatbot, manually tested by the human. "
    "This test is documented separately and is not counted as Harness B. It was added to make the difference between "
    "AI-generated content and direct filesystem execution clearer, and to complete the manual testing using plain GPT."
)

# ==================== SECTION 3 ====================
doc.add_heading("3. Test 1 — Simple Task (Both Harnesses)", level=1)
doc.add_paragraph("Task: create hello.txt and notes.txt in a test folder.")

doc.add_heading("Claude Code", level=2)
doc.add_paragraph(
    "Claude Code created the test folder, wrote both files, and verified them. The result was real files on disk in "
    "project5-output/. This demonstrated direct filesystem execution — the AI harness autonomously created and wrote files."
)

doc.add_heading("VS Code GPT Extension", level=2)
doc.add_paragraph(
    "The VS Code GPT extension was given the same task with an explicit instruction to actually create the files. "
    "It provided GPT-style assistance but did not create the files itself. A clean retest left chatgpt-test-output/ empty, "
    "confirming that no autonomous filesystem creation was observed. The human remained the execution gate."
)

# ==================== SECTION 4 (NEW) ====================
doc.add_heading("4. Test 1A — Plain ChatGPT Manual Reference (Completed Using Plain GPT)", level=1)

doc.add_paragraph(
    "Plain ChatGPT was given the same type of low-stakes task: create hello.txt and notes.txt with a greeting and the current date."
)

doc.add_paragraph(
    "ChatGPT generated the requested content in the conversation but did not directly create the files on the Windows filesystem. "
    "The human manually created hello.txt and notes.txt in chatgpt-test-output/ using the generated content."
)

doc.add_paragraph(
    "This test was not presented as proof that ChatGPT created the files. The files existing afterward are the result of the "
    "human's manual filesystem action. The purpose was to distinguish conversational content generation from direct filesystem execution — "
    "and to document that even when plain ChatGPT was used, it did not demonstrate the autonomous Reach, Loop, or Gate capabilities "
    "that Claude Code showed."
)

doc.add_paragraph(
    "The file chatgpt-test-output/observation.md records the complete six-part observation for this manual test, confirming:"
)

doc.add_paragraph(
    "- ChatGPT generated the requested content in the conversation.",
    style="List Bullet"
)
doc.add_paragraph(
    "- ChatGPT did not create the files on disk — the human did.",
    style="List Bullet"
)
doc.add_paragraph(
    "- This demonstrates the difference between generating content and directly executing filesystem actions.",
    style="List Bullet"
)

doc.add_paragraph(
    "Key takeaway: Even with plain GPT, the pattern is identical to the VS Code GPT extension — content generation without autonomous filesystem reach."
)

# ==================== SECTION 5 ====================
doc.add_heading("5. Test 2 — Richer Task, Claude Code Only (Plan → Execute → Verify → Correct)", level=1)

doc.add_paragraph(
    "Test 1 exercised a simple straight-line pass. Test 2 was added specifically to obtain stronger evidence for Loop and Spine "
    "using a task with four real stages."
)

doc.add_paragraph(
    "Recorded tool sequence: mkdir → date → write task-plan.md → write hello.txt → date → write notes.md → "
    "read all three files back → verify against the plan."
)

doc.add_paragraph(
    "Loop result: plan, execute, and verify were genuinely exercised. The correction step was not exercised because nothing was "
    "wrong on the first pass. This remains an honest gap rather than a claimed capability."
)

doc.add_paragraph(
    "Spine result: proven through task state. notes.md correctly described the order of the earlier file writes, without needing "
    "to reconstruct that order by rereading the disk — proving that running task state was maintained across tool calls."
)

# ==================== SECTION 6 ====================
doc.add_heading("6. Six-Part Shape — Official Two-Harness Comparison", level=1)

table = doc.add_table(rows=1, cols=3)
table.style = "Table Grid"
table.alignment = WD_TABLE_ALIGNMENT.CENTER

# Header
header_cells = table.rows[0].cells
header_cells[0].text = "Part"
header_cells[1].text = "Claude Code"
header_cells[2].text = "VS Code GPT Extension"

# Data rows
rows_data = [
    ("Heartbeat", "Discrete progress/results associated with real tool calls.", "A single discrete GPT response."),
    ("Reach", "Direct filesystem access; created and read real files.", "No direct filesystem creation observed."),
    ("Loop", "Create → write → verify ran autonomously end to end.", "No autonomous filesystem loop; human created the files."),
    ("Spine", "Proven in Test 2 through maintained task state across tool calls.", "Kept conversation context, but no filesystem sequence to test."),
    ("Gate", "Available tool permissions allowed execution to proceed.", "Human remained the required gate for file creation."),
    ("Body", "Agent plus filesystem tools such as Bash, Write, and Read.", "GPT code-generation utility.")
]

for row_data in rows_data:
    row_cells = table.add_row().cells
    for i, text in enumerate(row_data):
        row_cells[i].text = text

doc.add_paragraph()

# ==================== SECTION 7 (NEW) ====================
doc.add_heading("7. Plain ChatGPT — Six-Part Manual Reference (Completed Using Plain GPT)", level=1)

doc.add_paragraph(
    "This section documents the six-part observation for the plain ChatGPT manual test. "
    "The test was performed separately using plain GPT (not the VS Code extension or Claude Code) "
    "to complete the same task manually."
)

table2 = doc.add_table(rows=1, cols=2)
table2.style = "Table Grid"
table2.alignment = WD_TABLE_ALIGNMENT.CENTER

# Header
header2 = table2.rows[0].cells
header2[0].text = "Part"
header2[1].text = "Observed Behavior (Plain ChatGPT)"

# Data rows
plain_rows = [
    ("Heartbeat", "ChatGPT responded through the conversation with generated content."),
    ("Reach", "No direct filesystem reach was demonstrated — ChatGPT did not create files on disk."),
    ("Loop", "No autonomous create → write → verify filesystem loop occurred."),
    ("Spine", "The task and instructions were maintained through the conversation context."),
    ("Gate", "The human was the execution gate for creating the local files manually."),
    ("Body", "ChatGPT provided the requested text/content; the human performed the filesystem operation.")
]

for row_data in plain_rows:
    row_cells = table2.add_row().cells
    row_cells[0].text = row_data[0]
    row_cells[1].text = row_data[1]

doc.add_paragraph()

doc.add_paragraph(
    "This confirms that plain ChatGPT, like the VS Code GPT extension, generates assistance content "
    "but does not autonomously reach the filesystem. The key difference between plain ChatGPT and Claude Code "
    "is the presence of Reach, Loop, and Gate in the latter."
)

# ==================== SECTION 8 ====================
doc.add_heading("8. Three Required Implementation Differences", level=1)

for heading, text in [
    ("Reach", "Claude Code directly touched the filesystem; the VS Code GPT extension and plain ChatGPT did not demonstrate direct filesystem creation."),
    ("Loop", "Claude Code completed an autonomous create → write → verify sequence; the extension and plain ChatGPT stopped at generating assistance."),
    ("Gate", "Claude Code's available tool permissions allowed the task to proceed; with the extension and plain ChatGPT, the human was required to create the files.")
]:
    p = doc.add_paragraph()
    r = p.add_run(heading + " — ")
    r.bold = True
    p.add_run(text)

doc.add_paragraph()

# ==================== SECTION 9 ====================
doc.add_heading("9. Why The Simple Task Was Important", level=1)

doc.add_paragraph(
    "Creating two small text files is not special because of the files themselves. The task was intentionally simple so that the "
    "harness capabilities could be isolated. A complicated task could fail for many reasons; a simple task gives a cleaner signal "
    "about whether the harness can reach the filesystem, execute a sequence, verify the result, and leave the human as a gate or not."
)

doc.add_paragraph(
    "The meaningful observation is therefore not 'Claude Code can write hello.txt.' It is that Claude Code demonstrated a path from "
    "AI instruction to real tool execution and filesystem state, while the other tested conversational/code-generation harnesses did not "
    "demonstrate that same direct reach — even when using plain GPT manually."
)

# ==================== SECTION 10 ====================
doc.add_heading("10. Folders and Files Involved", level=1)

folders_list = [
    "project5-output/ — Test 1 real output created by Claude Code.",
    "chatgpt-test-output/ — Test 1A plain ChatGPT manual test output.",
    "chatgpt-test-output/hello.txt — manually created using ChatGPT-generated content.",
    "chatgpt-test-output/notes.txt — manually created using ChatGPT-generated content.",
    "chatgpt-test-output/observation.md — six-part observation of the plain ChatGPT manual test.",
    "claude-code-test/agent-test-output/ — Test 2 real output: task-plan.md, hello.txt, notes.md.",
    "claude-code-test/agent-test-observations.md — detailed Test 2 observation.",
    "claude-code-observations.md — detailed Test 1 Claude Code observation.",
    "comparison.md — formal side-by-side comparison and implementation differences.",
    "README.md — project overview and structure.",
    "assignment.md — original Project 5 assignment."
]

for item in folders_list:
    doc.add_paragraph(item, style="List Bullet")

doc.add_paragraph()

# ==================== SECTION 11 ====================
doc.add_heading("11. Honest Limitation (Left Open On Purpose)", level=1)

doc.add_paragraph(
    "The correction step of plan → execute → verify → correct was not actually tested because the first pass had nothing wrong to fix. "
    "Therefore self-correction is not claimed as proven. A future test could deliberately introduce an incorrect value and observe whether "
    "the agent detects and repairs it."
)

# ==================== SECTION 12 ====================
doc.add_heading("12. Final Verdict", level=1)

p = doc.add_paragraph()
r = p.add_run("PROJECT 5 — COMPLETE")
r.bold = True
r.font.size = Pt(14)

doc.add_paragraph(
    "The official two-harness comparison identifies three real implementation differences—Reach, Loop, and Gate—and maps the behavior "
    "to Heartbeat, Reach, Loop, Spine, Gate, and Body. Test 2 adds genuine, proven evidence for Spine and partial evidence for Loop, "
    "while explicitly and honestly flagging the one part (self-correction) that remains untested rather than assuming it works."
)

doc.add_paragraph(
    "The additional plain ChatGPT manual test provides supporting evidence showing that files created manually from ChatGPT-generated "
    "content should not be mistaken for direct AI filesystem execution. The pattern is clear: plain ChatGPT and the VS Code GPT extension "
    "both generate assistance content, but only Claude Code demonstrated autonomous Reach, Loop, and Gate capabilities."
)

# ==================== SAVE FILE ====================
# Save to current directory with a descriptive name
output_filename = "Project5_Updated_Full_Record_v2.docx"
doc.save(output_filename)
print(f"✅ Document created successfully: {output_filename}")
print(f"📁 Location: {os.path.abspath(output_filename)}")