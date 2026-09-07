from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = Document()

# Page setup
sec = doc.sections[0]
sec.top_margin = Inches(0.7)
sec.bottom_margin = Inches(0.7)
sec.left_margin = Inches(0.8)
sec.right_margin = Inches(0.8)

# Styles
styles = doc.styles
styles["Normal"].font.name = "Calibri"
styles["Normal"].font.size = Pt(10.5)
styles["Title"].font.name = "Calibri"
styles["Title"].font.size = Pt(20)
styles["Title"].font.bold = True
styles["Heading 1"].font.size = Pt(15)
styles["Heading 2"].font.size = Pt(12)

# Title
title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("Project 5 — Comparative Agent Harness Study")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Claude Code vs VS Code GPT")
r.bold = True

# 1. What Was Project 5?
doc.add_heading("1. What Was Project 5?", level=1)
doc.add_paragraph(
    "Project 5 was a practical comparison exercise from the General Agents Web Crash Course. "
    "The goal was to give the same low-stakes task to two AI harnesses, observe how each one actually "
    "handled the task, and map the observations to six concepts: Heartbeat, Reach, Loop, Spine, Gate, and Body."
)
doc.add_paragraph(
    "The assignment specifically asked for a one-page comparison that ignored superficial interface/button "
    "differences and identified the common six-part shape underneath. Cowork and ChatGPT Work were marked optional, "
    "so this study used the available Claude Code harness and the installed VS Code GPT extension as the two practical harnesses."
)

# 2. What Was Asked to Compare?
doc.add_heading("2. What Was Asked to Compare?", level=1)
doc.add_paragraph(
    "The same simple filesystem task was used for both harnesses: create two text files in the Project 5 test area, "
    "then verify what happened. The important question was not simply whether GPT could generate text. "
    "It was whether the harness could actually execute the task, what resources it could reach, whether it could run "
    "a multi-step loop, where the human remained a gate, and how the six-part agent structure appeared in practice."
)

# 3. Harness A — Claude Code
doc.add_heading("3. Harness A — Claude Code", level=1)
doc.add_paragraph(
    "Claude Code was given the task of creating a test folder and two files. It directly performed filesystem operations "
    "inside the Project 5 directory."
)
for text in [
    "Created the project5-output test folder.",
    "Created hello.txt containing a short greeting.",
    "Created notes.txt containing the date.",
    "Verified the created files and their contents through filesystem commands.",
    "Completed the sequence autonomously: create → write → verify."
]:
    doc.add_paragraph(text, style="List Bullet")
doc.add_paragraph(
    "Purpose of the test: determine whether the harness had direct filesystem reach and could complete a small "
    "multi-step task rather than merely generate instructions or text."
)
doc.add_paragraph(
    "Result: the test succeeded. Claude Code demonstrated direct filesystem reach and an observable execution loop."
)

# 4. Harness B — VS Code GPT Extension
doc.add_heading("4. Harness B — VS Code GPT Extension", level=1)
doc.add_paragraph(
    "The installed VS Code GPT extension (vscode-chat-gpt by ikasann-self) was given the same basic task, with an explicit "
    "instruction to actually create hello.txt and notes.txt in the test directory."
)
for text in [
    "The extension generated GPT-based assistance/output.",
    "It did not directly create the requested files in the local filesystem.",
    "In an earlier attempt, the human manually created the files by copying and pasting the generated text into VS Code.",
    "That manual result was deleted before the clean test.",
    "The clean retest left chatgpt-test-output empty, confirming that the extension itself did not demonstrate direct filesystem creation."
]:
    doc.add_paragraph(text, style="List Bullet")
doc.add_paragraph(
    "Purpose of the test: distinguish GPT text/code generation from actual agentic execution and filesystem reach."
)
doc.add_paragraph(
    "Result: the extension successfully provided GPT assistance, but the filesystem task was not autonomously executed. "
    "The human remained responsible for creating the files."
)

# 5. Why the Comparison Was Important
doc.add_heading("5. Why the Comparison Was Important", level=1)
doc.add_paragraph(
    "The comparison showed that two AI-assisted environments can produce useful output while having very different execution capabilities. "
    "The key difference was the harness, not merely the underlying language model response. Claude Code had tools that could directly act on "
    "the filesystem, while the VS Code GPT extension tested here functioned primarily as a GPT code-generation utility."
)

# 6. Three Required Implementation Differences
doc.add_heading("6. Three Required Implementation Differences", level=1)
table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = "Table Grid"
hdr = table.rows[0].cells
hdr[0].text = "Difference"
hdr[1].text = "Claude Code"
hdr[2].text = "VS Code GPT"
rows = [
    ("Reach", "Directly created, wrote, and read project files.", "No direct filesystem creation was observed."),
    ("Loop", "Executed create → write → verify as one workflow.", "No autonomous filesystem loop; human had to perform the file operation."),
    ("Gate", "Available filesystem tools allowed execution to proceed.", "Human remained the execution gate for local file creation.")
]
for a,b,c in rows:
    cells = table.add_row().cells
    cells[0].text, cells[1].text, cells[2].text = a,b,c

# 7. Six-Part Framework
doc.add_heading("7. Six-Part Framework — What We Monitored", level=1)
framework = [
    ("Heartbeat", "How the harness signals that work is happening. Claude Code produced discrete progress/results through tool calls; VS Code GPT produced a response rather than an independently running execution cycle."),
    ("Reach", "What the harness can access or act upon. Claude Code reached the local filesystem; VS Code GPT did not demonstrate that capability."),
    ("Loop", "How execution proceeds through steps, feedback, verification, and iteration. Claude Code performed create → write → verify; VS Code GPT stopped at assistance/output."),
    ("Spine", "The structural sequence holding the task together. Claude Code maintained the sequence across tool calls; VS Code GPT maintained request/response context without a filesystem execution sequence."),
    ("Gate", "Where permission or human control determines whether an action occurs. Claude Code used its available tool permissions; VS Code GPT required the human to perform the local file operation."),
    ("Body", "The actual implementation that performs the work. Claude Code used an agent plus filesystem tools such as Bash/Write; VS Code GPT acted as a GPT code-generation utility.")
]
for name, desc in framework:
    p = doc.add_paragraph()
    run = p.add_run(name + " — ")
    run.bold = True
    p.add_run(desc)

# 8. Important Folders and Files Used
doc.add_heading("8. Important Folders and Files Used", level=1)
items = [
    "project5-output/ — successful Claude Code test output.",
    "project5-output/hello.txt — greeting created by Claude Code.",
    "project5-output/notes.txt — date file created by Claude Code.",
    "chatgpt-test-output/ — clean VS Code GPT test directory; it remained empty after the direct-creation test.",
    "claude-code-observations.md — detailed Claude Code observation and six-part mapping.",
    "observations/chatgpt-work-run.md — VS Code GPT observation.",
    "comparison.md — final side-by-side comparison and the three implementation differences.",
    "README.md — project overview, harnesses, concepts, and structure.",
    "assignment.md — original Project 5 assignment.",
    "EXECUTION_GUIDE.md — execution/testing guidance.",
    "observations/framework.md — framework reference used for mapping observations."
]
for item in items:
    doc.add_paragraph(item, style="List Bullet")

# 9. What the Test Ultimately Demonstrated
doc.add_heading("9. What the Test Ultimately Demonstrated", level=1)
doc.add_paragraph(
    "The test was successful because it exposed a meaningful implementation difference while preserving the same conceptual six-part shape. "
    "Claude Code behaved as an executable agent harness with direct filesystem tools. VS Code GPT behaved as an AI assistance/code-generation "
    "harness in this test, with the human performing the filesystem action."
)
doc.add_paragraph(
    "The main lesson of Project 5 is therefore: do not judge an agent system only by its interface or generated text. "
    "Observe its heartbeat, reach, loop, spine, gate, and body to understand what it can actually do."
)

# 10. Final Project 5 Verdict
doc.add_heading("10. Final Project 5 Verdict", level=1)
p = doc.add_paragraph()
r = p.add_run("PROJECT 5 — COMPLETE")
r.bold = True
doc.add_paragraph(
    "The comparison contains the required three implementation differences (Reach, Loop, Gate) and maps the observed behavior "
    "to all six concepts: Heartbeat, Reach, Loop, Spine, Gate, and Body."
)

# Save
doc.save("Project_5_Summary.docx")
print("✅ Document saved as: Project_5_Summary.docx")